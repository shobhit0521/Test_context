"""Builds the presentation deck: ContextAI_Evaluation.pptx.

Covers the whole project across all three repos - ContextAI (the code-graph
engine), ConnectContext (the MCP bridge that exposes it to LLM agents), and
Test_context (this repo, the evaluation that measures whether the bridge
actually helps a real coding agent) - in a dense, 12-slide deck.

Run make_diagrams.py and make_charts.py first (or via `python3 build_deck.py`,
which does it automatically) to (re)generate the image assets this script
places on slides.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT_PATH = HERE / "ContextAI_Evaluation.pptx"

# ---- theme ----
NAVY = RGBColor(0x1A, 0x1A, 0x2E)
GREEN = RGBColor(0x2C, 0xA0, 0x2C)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
LIGHT_BG = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
AMBER = RGBColor(0xD9, 0x77, 0x06)
BLUE = RGBColor(0x25, 0x63, 0xEB)
PURPLE = RGBColor(0x93, 0x33, 0xEA)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT = "Calibri"


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def _fill_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def _add_text(slide, x, y, w, h, text, size=18, bold=False, color=DARK_TEXT,
              align=PP_ALIGN.LEFT, italic=False, font=FONT, line_spacing=1.12,
              anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = line_spacing
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.name = font
            run.font.color.rgb = color
    return box


def _add_bullets(slide, x, y, w, h, items, size=17, color=DARK_TEXT,
                  bullet_color=GREEN, spacing_after=8, font=FONT, line_spacing=1.05):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing_after)
        p.line_spacing = line_spacing
        run_bullet = p.add_run()
        run_bullet.text = "\u25B8  "
        run_bullet.font.size = Pt(size)
        run_bullet.font.color.rgb = bullet_color
        run_bullet.font.bold = True
        run_bullet.font.name = font
        run_text = p.add_run()
        run_text.text = item
        run_text.font.size = Pt(size)
        run_text.font.color.rgb = color
        run_text.font.name = font
    return box


def _accent_bar(slide, color=GREEN, y=Inches(1.22), height=Pt(4), x=Inches(0.55), w=Inches(2.2)):
    bar = slide.shapes.add_shape(1, x, y, w, height)  # MSO_SHAPE.RECTANGLE = 1
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_title_bar(slide, title, subtitle=None, title_color=NAVY, accent=GREEN, kicker=None):
    top = Inches(0.32)
    if kicker:
        _add_text(slide, Inches(0.55), top, Inches(12.2), Inches(0.32), kicker,
                  size=13, bold=True, color=accent)
        top = Inches(0.62)
    _add_text(slide, Inches(0.55), top, Inches(12.2), Inches(0.75), title,
              size=29, bold=True, color=title_color)
    _accent_bar(slide, color=accent, y=top + Inches(0.72))
    if subtitle:
        _add_text(slide, Inches(0.55), top + Inches(0.82), Inches(12.2), Inches(0.4), subtitle,
                   size=14, color=RGBColor(0x66, 0x66, 0x66), italic=True)


def add_content_slide(prs, title, subtitle=None, bg=WHITE, accent=GREEN, kicker=None):
    slide = _blank_slide(prs)
    _fill_bg(slide, bg)
    add_title_bar(slide, title, subtitle, accent=accent, kicker=kicker)
    return slide


def add_picture_fit(slide, img_path, x, y, max_w, max_h):
    from PIL import Image

    with Image.open(img_path) as im:
        iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    px = x + (max_w - w) // 2
    py = y + (max_h - h) // 2
    slide.shapes.add_picture(str(img_path), px, py, width=w, height=h)


def add_footer(slide, text, page_no=None):
    _add_text(slide, Inches(0.55), Inches(7.12), Inches(9), Inches(0.32), text,
              size=10, color=RGBColor(0xAA, 0xAA, 0xAA))
    if page_no is not None:
        _add_text(slide, Inches(12.5), Inches(7.12), Inches(0.6), Inches(0.32), str(page_no),
                  size=10, color=RGBColor(0xAA, 0xAA, 0xAA), align=PP_ALIGN.RIGHT)


def _callout_box(slide, x, y, w, h, text, fill=RGBColor(0xEE, 0xF6, 0xEE), border=GREEN, size=22):
    box = slide.shapes.add_shape(1, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = border
    box.line.width = Pt(2.5)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = NAVY
    run.font.name = FONT
    return box


# --------------------------------------------------------------------------
# Slide builders — 12 slides total, each dense with content.
# --------------------------------------------------------------------------

def slide_title(prs):
    slide = _blank_slide(prs)
    _fill_bg(slide, NAVY)

    _add_text(slide, Inches(1.0), Inches(1.35), Inches(11.3), Inches(1.9),
              "Does Giving an AI a \u201cMap\u201d of Your Code\nActually Help?",
              size=38, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    band = slide.shapes.add_shape(1, Inches(0), Inches(3.55), SLIDE_W, Inches(0.06))
    band.fill.solid()
    band.fill.fore_color.rgb = GREEN
    band.line.fill.background()

    _add_text(slide, Inches(1.5), Inches(3.8), Inches(10.3), Inches(0.55),
              "The whole project, in one deck: three repos, one pipeline, one real answer",
              size=18, color=GRAY, align=PP_ALIGN.CENTER, italic=True)

    # three repo chips
    chips = [
        ("ContextAI", "code-graph engine", GREEN),
        ("ConnectContext", "MCP bridge", BLUE),
        ("Test_context", "the evaluation", AMBER),
    ]
    chip_w, chip_h, gap = Inches(3.2), Inches(1.0), Inches(0.35)
    total_w = chip_w * 3 + gap * 2
    start_x = (SLIDE_W - total_w) // 2
    y = Inches(4.6)
    for i, (name, desc, color) in enumerate(chips):
        x = start_x + i * (chip_w + gap)
        chip = slide.shapes.add_shape(1, x, y, chip_w, chip_h)
        chip.fill.solid()
        chip.fill.fore_color.rgb = RGBColor(0x24, 0x24, 0x3E)
        chip.line.color.rgb = color
        chip.line.width = Pt(1.75)
        tf = chip.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r1 = p.add_run()
        r1.text = name
        r1.font.size = Pt(16)
        r1.font.bold = True
        r1.font.color.rgb = WHITE
        r1.font.name = FONT
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = desc
        r2.font.size = Pt(11.5)
        r2.font.italic = True
        r2.font.color.rgb = color
        r2.font.name = FONT

    _add_text(slide, Inches(1.5), Inches(6.15), Inches(10.3), Inches(0.5),
              "48 real questions \u00b7 4 real codebases \u00b7 288 real AI runs",
              size=14, color=RGBColor(0x88, 0x88, 0x99), align=PP_ALIGN.CENTER)
    return slide


def slide_problem_and_idea(prs):
    slide = add_content_slide(prs, "The problem, the idea, and the question we tested")
    _add_bullets(slide, Inches(0.55), Inches(1.85), Inches(6.4), Inches(3.0), [
        "Real codebases are huge. Asking an AI \u201cwho calls this?\u201d normally means it greps and opens files one at a time \u2014 slow, and easy to miss something in a file never opened.",
        "The idea: pre-build a \u201cmap\u201d of the codebase \u2014 every function, every connection \u2014 once, and hand it to the AI as an extra tool instead of making it re-derive that structure every time.",
        "That map-building idea is exactly what ContextAI + ConnectContext (the first two repos below) set out to build.",
    ], size=15.5, spacing_after=10)
    _callout_box(slide, Inches(0.55), Inches(5.05), Inches(6.4), Inches(1.55),
                 "\u201cDoes a real AI coding agent answer code questions better with the map \u2014 vs. its normal tools alone?\u201d",
                 fill=LIGHT_BG, border=GREEN, size=16.5)
    add_picture_fit(slide, ASSETS / "search_vs_map.png", Inches(7.15), Inches(1.85), Inches(5.6), Inches(4.85))
    add_footer(slide, "Part 1 \u2014 The Idea")
    return slide


def slide_three_repos(prs):
    slide = add_content_slide(prs, "The whole project spans three repos", accent=BLUE)
    add_picture_fit(slide, ASSETS / "three_repos.png", Inches(0.35), Inches(1.75), Inches(12.6), Inches(5.15))
    add_footer(slide, "Part 2 \u2014 The Whole Project")
    return slide


def slide_contextai(prs):
    slide = add_content_slide(prs, "Repo 1 \u2014 ContextAI: the code-graph engine", accent=GREEN,
                               kicker="github.com/shobhit0521/ContextAI")
    add_picture_fit(slide, ASSETS / "contextai_arch.png", Inches(0.35), Inches(1.7), Inches(12.6), Inches(5.25))
    add_footer(slide, "Part 2 \u2014 The Whole Project")
    return slide


def slide_connectcontext(prs):
    slide = add_content_slide(prs, "Repo 2 \u2014 ConnectContext: the MCP bridge", accent=BLUE,
                               kicker="github.com/shobhit0521/ConnectContext \u00b7 pypi.org/project/contextai-mcp")
    add_picture_fit(slide, ASSETS / "connectcontext_tools.png", Inches(0.35), Inches(1.7), Inches(12.6), Inches(5.3))
    add_footer(slide, "Part 2 \u2014 The Whole Project")
    return slide


def slide_methodology(prs):
    slide = add_content_slide(prs, "Repo 3 \u2014 Test_context: how we evaluated it", accent=AMBER,
                               kicker="github.com/shobhit0521/Test_context (this repo)")
    add_picture_fit(slide, ASSETS / "methodology.png", Inches(0.35), Inches(1.7), Inches(12.6), Inches(5.3))
    add_footer(slide, "Part 3 \u2014 Testing It For Real")
    return slide


def slide_hiccups(prs):
    slide = add_content_slide(prs, "Two real engineering hiccups, across two repos", accent=AMBER)
    _add_text(slide, Inches(0.6), Inches(1.85), Inches(5.85), Inches(0.5), "\U0001F41B  ConnectContext \u00d7 ContextAI: a hidden naming clash",
              size=17.5, bold=True, color=AMBER)
    _add_bullets(slide, Inches(0.6), Inches(2.4), Inches(5.85), Inches(4.0), [
        "ConnectContext launches as \u201cpython -m contextai_mcp,\u201d which adds the current directory to Python's import path.",
        "One target project (Flask) ships its own file named typing.py \u2014 same name as the Python stdlib module.",
        "Run with cwd inside Flask's source tree, that silently hijacked the real typing import and crashed the server on startup.",
        "Fixed by pinning ConnectContext's own cwd to a safe folder, away from whatever project it's mapping.",
    ], size=14.5, spacing_after=9)

    _add_text(slide, Inches(6.9), Inches(1.85), Inches(5.85), Inches(0.5), "\u23F1\uFE0F  Test_context \u00d7 Codex: a permissions limitation",
              size=17.5, bold=True, color=AMBER)
    _add_bullets(slide, Inches(6.9), Inches(2.4), Inches(5.85), Inches(4.0), [
        "Codex's unattended mode auto-cancels MCP tool-call approvals instead of granting them \u2014 a known open Codex CLI issue.",
        "Left as default, Arm B could never actually call ConnectContext's tools when run unattended.",
        "Fixed with a bypass flag \u2014 applied identically to both Arm A and Arm B, so tool availability is the only difference, not sandboxing.",
        "Both hiccups were root-caused and fixed by tracing through all three repos, not just the one that failed.",
    ], size=14.5, spacing_after=9)
    add_footer(slide, "Part 3 \u2014 Testing It For Real")
    return slide


def slide_results(prs):
    slide = add_content_slide(prs, "The headline result: roughly a coin flip", accent=GREEN)
    add_picture_fit(slide, ASSETS / "chart_headline.png", Inches(0.4), Inches(1.75), Inches(6.5), Inches(4.55))
    add_picture_fit(slide, ASSETS / "chart_by_project.png", Inches(6.95), Inches(1.75), Inches(6.1), Inches(4.55))
    _add_text(slide, Inches(0.55), Inches(6.4), Inches(12.2), Inches(0.65),
              "288 total AI runs, zero crashes. No project favored the map \u2014 not even Rich, the largest/most complex codebase tested.",
              size=13.5, color=RGBColor(0x66, 0x66, 0x66), italic=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Part 4 \u2014 Results")
    return slide


def slide_results_cost(prs):
    slide = add_content_slide(prs, "The map wasn't free \u2014 and wasn't more accurate either", accent=AMBER)
    add_picture_fit(slide, ASSETS / "chart_cost.png", Inches(0.4), Inches(1.75), Inches(6.5), Inches(4.6))
    add_picture_fit(slide, ASSETS / "chart_recall_gap.png", Inches(6.95), Inches(1.9), Inches(6.1), Inches(4.3))
    _add_text(slide, Inches(0.55), Inches(6.45), Inches(12.2), Inches(0.6),
              "Objective F1 was a near-tie too: 0.740 (no map) vs. 0.745 (with map) across 87 checkable answers.",
              size=13.5, color=RGBColor(0x66, 0x66, 0x66), italic=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Part 4 \u2014 Results")
    return slide


def slide_why(prs):
    slide = add_content_slide(prs, "Why didn't it help more?", accent=BLUE)
    _add_text(slide, Inches(0.6), Inches(1.8), Inches(5.9), Inches(0.5),
              "1. The referee was sometimes wrong", size=18, bold=True, color=BLUE)
    _add_bullets(slide, Inches(0.6), Inches(2.35), Inches(5.9), Inches(3.2), [
        "We hand-checked one \u201closs\u201d for the map: the AI called a buffer \u201cthread-local\u201d and the judge marked it wrong.",
        "Rich's actual source confirms the AI was right \u2014 the judge made the mistake.",
        "With 58/144 judged queries already ties, the true gap is likely smaller than the raw 45.1% vs. 54.9% split suggests.",
    ], size=14.5, spacing_after=10)

    _add_text(slide, Inches(6.9), Inches(1.8), Inches(5.85), Inches(0.5),
              "2. Trusting the map too much", size=18, bold=True, color=AMBER)
    _add_bullets(slide, Inches(6.9), Inches(2.35), Inches(5.85), Inches(3.2), [
        "The map is honest about its own gaps \u2014 list_gaps flags exactly what static analysis couldn't resolve.",
        "But when the AI treated the map's answer as complete instead of double-checking, it sometimes missed real connections a plain search would catch.",
        "Signature: when one arm's answer was a strict subset of the other's, it was the map-assisted arm 19 times out of 20.",
    ], size=14.5, spacing_after=10)

    _callout_box(slide, Inches(1.3), Inches(5.75), Inches(10.7), Inches(1.15),
                 "Bottom line: giving today's agent an extra map tool didn't clearly make it better \u2014 it used the tool eagerly, but that didn't translate into better answers.",
                 fill=LIGHT_BG, border=GREEN, size=15.5)
    add_footer(slide, "Part 5 \u2014 Why, and What's Next")
    return slide


def slide_reflection_and_forward(prs):
    slide = _blank_slide(prs)
    _fill_bg(slide, NAVY)
    _add_text(slide, Inches(0.6), Inches(0.4), Inches(11.5), Inches(0.65),
              "A personal note, and what this means going forward", size=25, bold=True, color=WHITE)
    _accent_bar(slide, color=AMBER, y=Inches(1.1), x=Inches(0.6))

    _add_text(slide, Inches(0.6), Inches(1.4), Inches(11.9), Inches(1.75),
              "When this problem first came up a few months ago, it felt urgent: AI models couldn't hold a big "
              "codebase \u201cin their head,\u201d so a pre-built map felt essential. In just those few months, AI models "
              "have gotten dramatically better at holding and reasoning over more code at once \u2014 that original "
              "problem has largely shrunk on its own, just from base models improving.",
              size=15.5, color=RGBColor(0xDD, 0xDD, 0xE5))

    _add_text(slide, Inches(0.6), Inches(3.35), Inches(11.9), Inches(0.5),
              "What this means going forward:", size=16, bold=True, color=GREEN)
    _add_bullets(slide, Inches(0.6), Inches(3.9), Inches(11.9), Inches(2.9), [
        "Tools like ContextAI/ConnectContext still have a place \u2014 but the bar for \u201cworth adding\u201d keeps rising as base models improve.",
        "The AI should be nudged to double-check the map's answer, not trust it outright \u2014 especially for \u201clist everything\u201d questions.",
        "The token/reading cost of an extra tool is real (~1.7x more per question here) and should be weighed against the unclear benefit.",
        "Worth re-running periodically across all three repos as both the models and the graph/MCP tooling keep evolving.",
    ], size=14.5, color=RGBColor(0xDD, 0xDD, 0xE5), bullet_color=GREEN, spacing_after=9)
    add_footer(slide, "Part 5 \u2014 Why, and What's Next")
    return slide


def slide_thankyou(prs):
    slide = _blank_slide(prs)
    _fill_bg(slide, NAVY)
    _add_text(slide, Inches(1.0), Inches(2.55), Inches(11.3), Inches(1.0), "Thank you",
              size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    band = slide.shapes.add_shape(1, Inches(5.5), Inches(3.55), Inches(2.3), Inches(0.06))
    band.fill.solid()
    band.fill.fore_color.rgb = GREEN
    band.line.fill.background()
    _add_text(slide, Inches(1.0), Inches(3.85), Inches(11.3), Inches(1.4),
              "github.com/shobhit0521/ContextAI\ngithub.com/shobhit0521/ConnectContext\ngithub.com/shobhit0521/Test_context",
              size=15, color=GRAY, align=PP_ALIGN.CENTER, italic=True, line_spacing=1.3)
    _add_text(slide, Inches(1.0), Inches(5.35), Inches(11.3), Inches(0.5),
              "Full write-up, charts, and raw data: eval/report/REPORT.md",
              size=13, color=RGBColor(0x88, 0x88, 0x99), align=PP_ALIGN.CENTER)
    return slide


def main():
    subprocess.run([sys.executable, str(HERE / "make_diagrams.py")], check=True)
    subprocess.run([sys.executable, str(HERE / "make_charts.py")], check=True)

    prs = new_presentation()

    slide_title(prs)
    slide_problem_and_idea(prs)
    slide_three_repos(prs)
    slide_contextai(prs)
    slide_connectcontext(prs)
    slide_methodology(prs)
    slide_hiccups(prs)
    slide_results(prs)
    slide_results_cost(prs)
    slide_why(prs)
    slide_reflection_and_forward(prs)
    slide_thankyou(prs)

    prs.save(OUT_PATH)
    print(f"Saved {OUT_PATH} ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
